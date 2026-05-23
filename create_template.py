import docx

doc = docx.Document("Review struktur Gedung RS Cokrodipo B. Lampung 2026.docx")
doc.add_heading("ETABS Automated Report Data", level=1)

doc.add_paragraph("Project Name: {{ project_name }}")
doc.add_paragraph("Engineer Name: {{ engineer_name }}")

doc.add_heading("Beam Forces", level=2)
# Adding a jinja tag table for docxtpl
table = doc.add_table(rows=2, cols=9)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ["Frame", "Station", "LoadCase", "P", "V2", "V3", "T", "M2", "M3"]
for i, header in enumerate(headers):
    hdr_cells[i].text = header

row_cells = table.rows[1].cells
row_cells[0].text = "{% tr for item in forces_list %}{{ item.Frame }}"
row_cells[1].text = "{{ item.Station }}"
row_cells[2].text = "{{ item.LoadCase }}"
row_cells[3].text = "{{ item.P }}"
row_cells[4].text = "{{ item.V2 }}"
row_cells[5].text = "{{ item.V3 }}"
row_cells[6].text = "{{ item.T }}"
row_cells[7].text = "{{ item.M2 }}"
row_cells[8].text = "{{ item.M3 }}{% endtr %}"

doc.add_heading("Beam Required Reinforcement", level=2)
table2 = doc.add_table(rows=2, cols=6)
table2.style = 'Table Grid'
hdr_cells2 = table2.rows[0].cells
headers2 = ["Frame", "Station", "TopRebar", "BotRebar", "ShearRebar", "TorsionRebar"]
for i, header in enumerate(headers2):
    hdr_cells2[i].text = header

row_cells2 = table2.rows[1].cells
row_cells2[0].text = "{% tr for item in rebar_list %}{{ item.Frame }}"
row_cells2[1].text = "{{ item.Station }}"
row_cells2[2].text = "{{ item.TopRebar }}"
row_cells2[3].text = "{{ item.BotRebar }}"
row_cells2[4].text = "{{ item.ShearRebar }}"
row_cells2[5].text = "{{ item.TorsionRebar }}{% endtr %}"

doc.save("Report_Template_Tagged.docx")
print("Template created successfully.")
